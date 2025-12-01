from jinja_excel_template.ExcelGenerator import ExcelGenerator
import os
from openpyxl import load_workbook

dir_path = os.path.dirname(os.path.realpath(__file__))


def test_generated_value_from_formula():
    xml = """
    <Excel>
      <sheet sheet_name="TestSheet">
        <row>
          <cell type="Number" formula="1+2" />
          <cell type="General" formula="CONCAT(['a','b'])" />
        </row>
      </sheet>
    </Excel>
    """

    out = os.path.join(dir_path, "temp", "formula_test1.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    # Default behavior: pass-through - formulas are stored
    wb = load_workbook(out, data_only=False)
    ws = wb["TestSheet"]
    assert isinstance(ws.cell(row=1, column=1).value, str)
    assert ws.cell(row=1, column=1).value.startswith("=")
    assert isinstance(ws.cell(row=1, column=2).value, str)
    assert ws.cell(row=1, column=2).value.startswith("=")


def test_generated_value_from_formula_evaluate_true():
    xml = """
    <Excel>
      <sheet sheet_name="TestSheetEval">
        <row>
          <cell type="Number" formula="1+2" evaluate="true" />
          <cell type="General" formula="CONCAT(['a','b'])" evaluate="true" />
        </row>
      </sheet>
    </Excel>
    """
    out = os.path.join(dir_path, "temp", "formula_test1_eval.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    wb = load_workbook(out, data_only=True)
    ws = wb["TestSheetEval"]
    assert ws.cell(row=1, column=1).value == 3
    assert ws.cell(row=1, column=2).value == 'ab'


def test_range_sum():
    xml = """
    <Excel>
      <sheet sheet_name="RangeSheet">
        <row>
          <cell type="Number" value="1" />
          <cell type="Number" value="2" />
          <cell type="Number" value="3" />
        </row>
        <row>
          <cell type="Number" formula="SUM(A1:C1)" />
        </row>
      </sheet>
    </Excel>
    """
    out = os.path.join(dir_path, "temp", "formula_test3.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    # Default is pass-through; formula string is saved
    wb = load_workbook(out, data_only=False)
    ws = wb["RangeSheet"]
    assert isinstance(ws.cell(row=2, column=1).value, str)
    assert ws.cell(row=2, column=1).value.startswith("=")


def test_range_sum_evaluate_true():
    xml = """
    <Excel>
      <sheet sheet_name="RangeSheet">
        <row>
          <cell type="Number" value="1" />
          <cell type="Number" value="2" />
          <cell type="Number" value="3" />
        </row>
        <row>
          <cell type="Number" formula="SUM(A1:C1)" evaluate="true" />
        </row>
      </sheet>
    </Excel>
    """
    out = os.path.join(dir_path, "temp", "formula_test3_eval.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    wb = load_workbook(out, data_only=True)
    ws = wb["RangeSheet"]
    assert ws.cell(row=2, column=1).value == 6


def test_cross_sheet_range_sum():
    xml = """
    <Excel>
      <sheet sheet_name="Sheet1">
        <row>
          <cell type="Number" value="4" />
          <cell type="Number" value="5" />
        </row>
      </sheet>
      <sheet sheet_name="Sheet2">
        <row>
          <cell type="Number" formula="SUM(Sheet1!A1:B1)" />
        </row>
      </sheet>
    </Excel>
    """
    out = os.path.join(dir_path, "temp", "formula_test4.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    wb = load_workbook(out, data_only=False)
    ws2 = wb["Sheet2"]
    assert isinstance(ws2.cell(row=1, column=1).value, str)
    assert ws2.cell(row=1, column=1).value.startswith("=")


def test_cross_sheet_range_sum_evaluate_true():
    xml = """
    <Excel>
      <sheet sheet_name="Sheet1">
        <row>
          <cell type="Number" value="4" />
          <cell type="Number" value="5" />
        </row>
      </sheet>
      <sheet sheet_name="Sheet2">
        <row>
          <cell type="Number" formula="SUM(Sheet1!A1:B1)" evaluate="true" />
        </row>
      </sheet>
    </Excel>
    """
    out = os.path.join(dir_path, "temp", "formula_test4_eval.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    wb = load_workbook(out, data_only=True)
    ws2 = wb["Sheet2"]
    assert ws2.cell(row=1, column=1).value == 9


def test_cross_sheet_get_single_cell():
    xml = """
    <Excel>
      <sheet sheet_name="First">
        <row>
          <cell type="Number" value="42" />
        </row>
      </sheet>
      <sheet sheet_name="Second">
        <row>
          <cell type="Number" formula="First!A1" />
        </row>
      </sheet>
    </Excel>
    """
    out = os.path.join(dir_path, "temp", "formula_test5.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    wb = load_workbook(out, data_only=False)
    assert wb["Second"].cell(row=1, column=1).value.startswith("=")

    # Evaluate server-side with evaluate=true
    xml2 = """
    <Excel>
      <sheet sheet_name="First">
        <row>
          <cell type="Number" value="42" />
        </row>
      </sheet>
      <sheet sheet_name="Second">
        <row>
          <cell type="Number" formula="First!A1" evaluate="true" />
        </row>
      </sheet>
    </Excel>
    """
    out2 = os.path.join(dir_path, "temp", "formula_test5_eval.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out2, xml_string=xml2)
    wb2 = load_workbook(out2, data_only=True)
    assert wb2["Second"].cell(row=1, column=1).value == 42


def test_circular_reference():
    xml = """
    <Excel>
      <sheet sheet_name="Circ">
        <row>
          <cell type="Number" formula="A2" />
        </row>
        <row>
          <cell type="Number" formula="A1" />
        </row>
      </sheet>
    </Excel>
    """
    out = os.path.join(dir_path, "temp", "formula_test6.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    wb = load_workbook(out, data_only=True)
    # Engine should detect circular or unresolved references and set error content
    assert wb["Circ"].cell(row=1, column=1).value == '#ERROR' or wb["Circ"].cell(row=1, column=1).value is None


def test_pass_through_formula_emitted():
    xml = """
    <Excel>
      <sheet sheet_name="TestSheet2">
        <row>
          <cell type="Number" formula="SUM([1,2,3])" pass_through="true" />
        </row>
      </sheet>
    </Excel>
    """

    out = os.path.join(dir_path, "temp", "formula_test2.xlsx")
    generator = ExcelGenerator()
    generator.generate_excel(out, xml_string=xml)
    wb = load_workbook(out, data_only=False)
    ws = wb["TestSheet2"]
    # the formula stored should start with '=' if configured
    assert ws.cell(row=1, column=1).value.startswith('=')
