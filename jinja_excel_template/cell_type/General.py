from openpyxl.worksheet.worksheet import Worksheet
from jinja_excel_template.helper.GetFont import get_font
from jinja_excel_template.helper.evaluate_formula import evaluate_formula


class General:
    def set_cell(self, worksheet: Worksheet, row: int, column: int, xml_cell):
        """
        This cell type set the cell as general cell. row and column are 1 based index.
        """
        if xml_cell is None:
            return

        formula = xml_cell.get("formula")
        evaluate = xml_cell.get("evaluate") or "false"  # evaluate server-side when true

        if formula:
            # Default: pass-through -> write Excel formula string
            if not (evaluate.lower() in ("true", "1")):
                # Ensure formula string starts with '=' for Excel
                formula_text = formula if formula.startswith("=") else f"={formula}"
                worksheet.cell(row, column, formula_text)
                return
            # If evaluate attribute is true, evaluate server-side and write the computed value
            try:
                # Build context mapping for A1 names and ranges
                ctx = {}
                # fill ctx with existing cell values
                max_col = worksheet.max_column or column
                max_row = row
                from openpyxl.utils import get_column_letter
                for r in range(1, max_row + 1):
                    for c in range(1, max_col + 1):
                        address = f"{get_column_letter(c)}{r}"
                        ctx[address] = worksheet.cell(r, c).value

                def range_resolver(start, end, sheet_name=None):
                    from openpyxl.utils import column_index_from_string
                    from openpyxl.utils.cell import coordinate_from_string
                    if sheet_name:
                        ws = worksheet.parent[sheet_name]
                    else:
                        ws = worksheet
                    sc, sr = coordinate_from_string(start)
                    ec, er = coordinate_from_string(end)
                    sc_i = column_index_from_string(sc)
                    ec_i = column_index_from_string(ec)
                    results = []
                    for rr in range(sr, er + 1):
                        for cc in range(sc_i, ec_i + 1):
                            addr = f"{get_column_letter(cc)}{rr}"
                            results.append(ws.cell(rr, cc).value)
                    return results

                ctx['__range_resolver__'] = range_resolver
                value = evaluate_formula(formula, context=ctx)
                if value is not None:
                    worksheet.cell(row, column, value)
                else:
                    worksheet.cell(row, column, None)
            except Exception:
                worksheet.cell(row, column, "#ERROR")
            return

        value = xml_cell.get("value")
        if value:
            worksheet.cell(row, column, value)

